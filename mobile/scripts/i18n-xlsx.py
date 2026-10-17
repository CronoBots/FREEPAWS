#!/usr/bin/env python3
"""Tableau de relecture des textes de l'app FreePaws (Excel), dans les deux sens.

  npm run i18n:export            -> i18n-relecture.xlsx (textes actuels + colonnes à remplir)
  npm run i18n:import -- <xlsx>  -> met à jour src/i18n/reviewed.json avec les textes finals

Les colonnes « Français final » et « English final » vides gardent le texte actuel.
Les variables entre accolades ({name}, {count}…) doivent rester identiques : sinon la ligne est refusée.
En français, les espaces insécables avant : ; ? ! » et après « sont remis automatiquement.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
REVIEWED = ROOT / "src" / "i18n" / "reviewed.json"
NBSP = " "

# Sections du dictionnaire -> (écran, public). Le français fait référence.
SECTIONS = {
    "common": ("Commun", "Clients"), "tabs": ("Onglets", "Clients"), "titles": ("Titres des écrans", "Clients"),
    "park": ("Parc (repris du site)", "Clients"), "status": ("Statut du parc", "Clients"),
    "camera": ("Caméra (repris du site)", "Clients"), "coaching": ("Coaching (repris du site)", "Clients"),
    "booking": ("Réservation", "Clients"), "account": ("Compte", "Clients"), "signIn": ("Connexion", "Clients"),
    "profile": ("Mes informations", "Clients"), "dogs": ("Mes chiens", "Clients"), "legal": ("Confidentialité", "Clients"),
    "notFound": ("Page introuvable", "Clients"), "config": ("Configuration", "Technique"),
    "errors": ("Messages d’erreur", "Clients"), "fiche": ("Fiches maître et chien", "Clients"),
    "parkBooking": ("Réservation du parc", "Clients"), "v11Client": ("Invités, questionnaire, urgence", "Clients"),
    "admin": ("Administration", "Administratrice"), "mfa": ("Double authentification", "Administratrice"),
    "adminClients": ("Administration — clients", "Administratrice"),
    "adminSafety": ("Administration — sécurité", "Administratrice"), "pricing": ("Administration — tarifs", "Administratrice"),
    "adminExtra": ("Administration — divers", "Administratrice"), "v11Admin": ("Administration — réservations", "Administratrice"),
}
HEADERS = ["Écran", "Public", "Clé (ne pas modifier)", "Français actuel", "Anglais actuel", "Français final", "English final", "Remarque"]
WIDTHS = [26, 14, 30, 60, 60, 60, 60, 30]


def dump():
    out = subprocess.run(["npx", "tsx", "scripts/i18n-dump.ts"], cwd=ROOT, check=True, capture_output=True, text=True)
    return json.loads(out.stdout)


def visible(text):
    # Excel affiche les espaces insécables comme des espaces : on les montre normalement.
    return text.replace(NBSP, " ")


def french_typography(text):
    text = re.sub(r"[  ]+([:;?!»])", NBSP + r"\1", text)
    text = re.sub(r"«[  ]+", "«" + NBSP, text)
    return text


def variables(text):
    return sorted(re.findall(r"\{\w+\}", text))


def export(path):
    wb = Workbook()
    guide = wb.active
    guide.title = "Mode d’emploi"
    for line in [
        "Relecture des textes de l’application FreePaws",
        "",
        "1. Onglet « Textes » : une ligne par texte affiché dans l’application.",
        "2. Écrivez le texte définitif dans « Français final » et/ou « English final ». Laissez vide pour garder le texte actuel.",
        "3. Gardez tels quels les mots entre accolades, par exemple {name} ou {count} : l’application les remplace (nom, nombre…).",
        "4. Les clés se terminant par _one / _other sont le singulier / pluriel du même texte.",
        "5. Ne modifiez pas la colonne « Clé ». Vous pouvez filtrer par « Écran » ou « Public ».",
        "6. Renvoyez le fichier : les textes finals seront intégrés à l’application.",
    ]:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Textes")
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2B3A30")
    for row in dump():
        section = row["key"].split(".")[0]
        screen, audience = SECTIONS.get(section, (section, "Clients"))
        ws.append([screen, audience, row["key"], visible(row["fr"]), visible(row["en"]), "", "", ""])
    final = PatternFill("solid", fgColor="FDFAEF")
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        r[5].fill = final
        r[6].fill = final
    for i, width in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = width
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    print(f"{ws.max_row - 1} textes exportés dans {path}")


def import_(path):
    current = {row["key"]: row for row in dump()}
    reviewed = json.loads(REVIEWED.read_text(encoding="utf-8"))
    ws = load_workbook(path, read_only=True)["Textes"]
    changed, refused = 0, []
    for values in ws.iter_rows(min_row=2, values_only=True):
        key = values[2]
        if not key or key not in current:
            continue
        for lang, column in (("fr", 5), ("en", 6)):
            text = values[column]
            if text is None or str(text).strip() == "":
                continue
            text = str(text).strip()
            if lang == "fr":
                text = french_typography(text)
            if variables(text) != variables(current[key]["fr"]):
                refused.append(f"{key} ({lang}) : variables {variables(text)} au lieu de {variables(current[key]['fr'])}")
                continue
            if text != current[key][lang]:
                reviewed[lang][key] = text
                changed += 1
    REVIEWED.write_text(json.dumps(reviewed, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(f"{changed} texte(s) mis à jour dans {REVIEWED.relative_to(ROOT)}")
    for line in refused:
        print("REFUSÉ :", line)
    return 1 if refused else 0


if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] not in {"export", "import"}:
        sys.exit(__doc__)
    if sys.argv[1] == "export":
        export(sys.argv[2] if len(sys.argv) > 2 else "i18n-relecture.xlsx")
    else:
        if len(sys.argv) < 3:
            sys.exit("Indiquez le fichier Excel relu.")
        sys.exit(import_(sys.argv[2]))
